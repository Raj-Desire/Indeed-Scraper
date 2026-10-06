"""
Match-accuracy evaluation harness.

1) Build a labelling sheet from past exported leads:
     python scripts/eval_matching.py bootstrap outputs/Indeed_Job_Leads_2026-07-21.xlsx outputs/Dice_Job_Leads_2026-09-22.xlsx --out eval/labels.csv
   Open eval/labels.csv and fill the `expected` column with High / Medium / Low
   (your sales team's judgement of real fit). Aim for 50-100 rows, mixed fit levels.

2) Score them with each engine and compare:
     python scripts/eval_matching.py run eval/labels.csv --engines v2,legacy

Reports band accuracy, High precision/recall, how often Low-fit jobs get flagged High,
mean error vs band midpoint, unscored rate, and a confusion matrix. Per-row results are
written to eval/results_<engine>.csv so disagreements can be inspected.
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.matching.eval_metrics import evaluate, normalize_label  # noqa: E402


def bootstrap(files: list[str], out: str) -> None:
    from openpyxl import load_workbook

    seen: set[str] = set()
    rows: list[dict] = []
    for f in files:
        try:
            ws = load_workbook(f, read_only=True, data_only=True).active
        except Exception as exc:
            print(f"skip {f}: {exc}")
            continue
        header_idx, headers = None, []
        data = list(ws.iter_rows(values_only=True))
        for i, r in enumerate(data):
            if r and "Job Title" in [str(c).strip() for c in r if c]:
                header_idx, headers = i, [str(c).strip() if c else "" for c in r]
                break
        if header_idx is None:
            print(f"skip {f}: no header row")
            continue
        col = {h: i for i, h in enumerate(headers)}
        jd_col = next((h for h in headers if h.lower() == "job description"), None)
        score_col = next((h for h in headers if "match" in h.lower() and "score" in h.lower()), None)
        for r in data[header_idx + 1:]:
            title = str(r[col["Job Title"]] or "").strip()
            jd = str(r[col[jd_col]] or "").strip() if jd_col else ""
            key = (title + jd[:200]).lower()
            if not title or len(jd) < 100 or key in seen:
                continue
            seen.add(key)
            rows.append({
                "title": title,
                "company": str(r[col["Company"]] or "") if "Company" in col else "",
                "job_description": jd,
                "previous_score": r[col[score_col]] if score_col else "",
                "expected": "",
            })
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["title", "company", "job_description", "previous_score", "expected"])
        w.writeheader()
        w.writerows(rows)
    print(f"Wrote {len(rows)} rows to {out}. Fill the 'expected' column (High/Medium/Low).")


async def _score(engine: str, jd: str, title: str, state: dict):
    if engine == "v2":
        r = await state["engine"].evaluate(jd, title)
    else:
        chunks = await state["kb"].search(jd)
        r = await state["matcher"].evaluate(jd, chunks)
    return r


async def run(csv_path: str, engines: list[str], out_dir: str) -> None:
    from app.knowledge_base.azure_search import AzureSearchKnowledgeBase
    from app.matching.engine import MatchEngine
    from app.matching.llm_matcher import LLMMatcher

    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        labelled = [r for r in csv.DictReader(fh) if normalize_label(r.get("expected", ""))]
    if not labelled:
        print("No rows with an 'expected' label found.")
        return

    kb, matcher = AzureSearchKnowledgeBase(), LLMMatcher()
    state = {"kb": kb, "matcher": matcher, "engine": MatchEngine(kb, matcher)}
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    summary = {}
    for eng in engines:
        results = []
        for i, r in enumerate(labelled, 1):
            res = await _score(eng, r["job_description"], r["title"], state)
            results.append({
                "title": r["title"], "expected": normalize_label(r["expected"]),
                "score": res.match_score, "status": res.match_status,
                "matched": "; ".join(res.matched_skills), "missing": "; ".join(res.missing_skills),
                "reason": res.match_reason,
            })
            print(f"[{eng}] {i}/{len(labelled)} {r['title'][:50]} -> {res.match_score}", flush=True)
        summary[eng] = evaluate(results)
        with open(Path(out_dir) / f"results_{eng}.csv", "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=list(results[0].keys()))
            w.writeheader()
            w.writerows(results)
    await kb.close()
    await matcher.close()
    print(json.dumps(summary, indent=2))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("bootstrap")
    b.add_argument("files", nargs="+")
    b.add_argument("--out", default="eval/labels.csv")
    r = sub.add_parser("run")
    r.add_argument("csv")
    r.add_argument("--engines", default="v2,legacy")
    r.add_argument("--out-dir", default="eval")
    a = ap.parse_args()
    if a.cmd == "bootstrap":
        bootstrap(a.files, a.out)
    else:
        asyncio.run(run(a.csv, [e.strip() for e in a.engines.split(",")], a.out_dir))


if __name__ == "__main__":
    main()
