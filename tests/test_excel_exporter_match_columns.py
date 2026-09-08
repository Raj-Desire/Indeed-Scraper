"""Confirms the Excel export includes the new match columns and keeps the Job URL hyperlink working."""
import shutil
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from app.excel.exporter import ExcelExporter
from app.models.job import JobPosting

_OUT_DIR = Path(__file__).resolve().parent / "_tmp_excel_out"


def test_export_includes_match_columns_and_url_hyperlink():
    if _OUT_DIR.exists():
        shutil.rmtree(_OUT_DIR)

    job = JobPosting(
        job_title="SharePoint Developer",
        company="Acme Corp",
        job_description="Need SPFx skills",
        job_url="https://www.indeed.com/viewjob?jk=abc123",
    )
    job.match_score = 88
    job.matched_skills = ["SharePoint", "SPFx"]
    job.missing_skills = ["Dynamics 365"]
    job.match_reason = "Strong SPFx overlap"

    path = ExcelExporter().export([job], output_dir=str(_OUT_DIR))
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    assert "Match Score" in headers
    assert "Matched Skills" in headers
    assert "Missing Skills" in headers
    assert "Match Reason" in headers
    assert "Job URL" in headers

    url_col = headers.index("Job URL") + 1
    url_cell = ws.cell(row=2, column=url_col)
    assert url_cell.hyperlink.target == "https://www.indeed.com/viewjob?jk=abc123"
    assert url_cell.value == "View on Indeed"

    match_score_col = headers.index("Match Score") + 1
    assert ws.cell(row=2, column=match_score_col).value == 88

    shutil.rmtree(_OUT_DIR)


if __name__ == "__main__":
    test_export_includes_match_columns_and_url_hyperlink()
    print("OK")
