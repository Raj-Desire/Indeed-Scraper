"""
Tests for Dice Excel Export Functionality
=========================================
Verifies that:
1. DiceService.export_excel() generates a styled workbook with Dice banner and 'View on Dice' links.
2. DiceService.export_excel() supports selected_ids filtering and captures search parameters.
3. GET /api/dice/export/excel and POST /api/dice/export/excel endpoints work properly.
4. Error handling when no leads exist returns HTTP 400.
"""

import shutil
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi import FastAPI
from fastapi.testclient import TestClient
import openpyxl
import pytest

from app.models.job import JobPosting
from app.services.dice_service import DiceService
import app.dashboard.router as router_mod

_OUT_DIR = Path(__file__).resolve().parent / "_tmp_dice_excel_out"


def setup_function():
    if _OUT_DIR.exists():
        shutil.rmtree(_OUT_DIR)
    _OUT_DIR.mkdir(parents=True, exist_ok=True)


def teardown_function():
    if _OUT_DIR.exists():
        shutil.rmtree(_OUT_DIR)


def test_dice_service_export_excel_formatting():
    service = DiceService()
    job1 = JobPosting(
        job_title="Senior Python Architect",
        company="Tech Innovators Inc",
        country="United States",
        location_remote_type="Remote",
        experience="5+ years",
        salary_range="$140,000 - $170,000",
        industry="Information Technology",
        company_size="100-500",
        job_description="Architect scalable microservices in Python and AWS.",
        job_url="https://www.dice.com/job-detail/abc-123",
        lead_source="Dice",
        job_summary="High priority role matching core Python & Cloud architecture.",
    )
    job1.match_score = 95
    job1.matched_skills = ["Python", "AWS", "Microservices"]
    job1.missing_skills = ["Kubernetes"]

    job2 = JobPosting(
        job_title="React Frontend Lead",
        company="Digital Solutions Ltd",
        country="Canada",
        location_remote_type="Remote",
        experience="4+ years",
        salary_range="CAD 120,000",
        industry="Software",
        company_size="50-200",
        job_description="Lead modern React and TypeScript frontend development.",
        job_url="https://www.dice.com/job-detail/def-456",
        lead_source="Dice",
        job_summary="Solid frontend lead opportunity.",
    )
    job2.match_score = 82
    job2.matched_skills = ["React", "TypeScript"]
    job2.missing_skills = ["Next.js"]

    service._results = [job1, job2]
    service._last_search_params = {
        "keywords": ["python", "react"],
        "countries": ["United States", "Canada"],
        "posted_date": "1",
        "workplace_types": ["Remote"],
    }

    excel_path = service.export_excel(output_dir=str(_OUT_DIR))
    assert excel_path.exists()
    assert "Dice_Job_Leads_" in excel_path.name

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    assert ws.title == "Job Leads"

    # Executive Banner
    banner_val = str(ws.cell(row=1, column=1).value)
    assert "DICE JOB SOURCING REPORT" in banner_val

    # Parameters Card
    assert ws.cell(row=3, column=1).value == "Selected Countries:"
    assert "United States" in str(ws.cell(row=3, column=2).value)
    assert ws.cell(row=3, column=6).value == "Date Posted Filter:"
    assert "Last 24 hours" in str(ws.cell(row=3, column=7).value)
    assert ws.cell(row=4, column=1).value == "Search Keyword / Role:"
    assert "python, react" in str(ws.cell(row=4, column=2).value)
    assert ws.cell(row=4, column=6).value == "Location Filter:"
    assert "Fully Remote Only" in str(ws.cell(row=4, column=7).value)
    assert ws.cell(row=5, column=6).value == "Total Leads:"
    assert "2 leads captured" in str(ws.cell(row=5, column=7).value)

    # Headers & Data Rows
    headers = [cell.value for cell in ws[7]]
    assert "#" in headers
    assert "Job Title" in headers
    assert "Match Score" in headers
    assert "Matched Skills" in headers
    assert "Job URL" in headers

    # First data row
    url_col = headers.index("Job URL") + 1
    cell_url = ws.cell(row=8, column=url_col)
    assert cell_url.value == "View on Dice"
    assert cell_url.hyperlink.target == "https://www.dice.com/job-detail/abc-123"

    score_col = headers.index("Match Score") + 1
    assert ws.cell(row=8, column=score_col).value == 95
    assert ws.cell(row=9, column=score_col).value == 82


def test_dice_service_export_excel_selected_ids_filter():
    service = DiceService()
    job1 = JobPosting(job_title="Dev 1", company="Acme 1", lead_source="Dice", job_url="https://dice.com/1")
    job2 = JobPosting(job_title="Dev 2", company="Acme 2", lead_source="Dice", job_url="https://dice.com/2")
    service._results = [job1, job2]

    excel_path = service.export_excel(selected_ids=[str(job2.id)], output_dir=str(_OUT_DIR))
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    # Total leads in header
    assert "1 leads captured" in str(ws.cell(row=5, column=7).value)
    # Row 8 is job2, row 9 is empty
    assert ws.cell(row=8, column=2).value == "Dev 2"
    assert ws.cell(row=9, column=2).value is None


def test_dice_service_export_excel_empty_raises():
    service = DiceService()
    service._results = []
    with pytest.raises(ValueError, match="No Dice job leads to export"):
        service.export_excel(output_dir=str(_OUT_DIR))


def _make_test_client(fake_service, monkeypatch):
    app = FastAPI()
    app.include_router(router_mod.router)
    monkeypatch.setattr(router_mod, "get_dice_service", lambda: fake_service)
    return TestClient(app)


def test_get_dice_export_excel_route(monkeypatch):
    service = DiceService()
    job = JobPosting(job_title="Cloud Engineer", company="CloudCorp", lead_source="Dice", job_url="https://dice.com/cloud")
    service._results = [job]

    client = _make_test_client(service, monkeypatch)
    resp = client.get("/api/dice/export/excel")
    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.headers.get("content-type", "")
    assert "Dice_Job_Leads_" in resp.headers.get("content-disposition", "")


def test_get_dice_export_excel_route_selected_ids(monkeypatch):
    service = DiceService()
    job1 = JobPosting(job_title="Cloud Engineer 1", company="CloudCorp 1", lead_source="Dice")
    job2 = JobPosting(job_title="Cloud Engineer 2", company="CloudCorp 2", lead_source="Dice")
    service._results = [job1, job2]

    client = _make_test_client(service, monkeypatch)
    resp = client.get(f"/api/dice/export/excel?selected_ids={job1.id}")
    assert resp.status_code == 200


def test_get_dice_export_excel_route_no_leads(monkeypatch):
    service = DiceService()
    service._results = []

    client = _make_test_client(service, monkeypatch)
    resp = client.get("/api/dice/export/excel")
    assert resp.status_code == 400
    assert "No Dice job leads to export" in resp.json()["detail"]


def test_post_dice_export_excel_route(monkeypatch):
    service = DiceService()
    job1 = JobPosting(job_title="Cloud Engineer 1", company="CloudCorp 1", lead_source="Dice")
    job2 = JobPosting(job_title="Cloud Engineer 2", company="CloudCorp 2", lead_source="Dice")
    service._results = [job1, job2]

    client = _make_test_client(service, monkeypatch)
    resp = client.post("/api/dice/export/excel", json={"selected_ids": [str(job2.id)]})
    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.headers.get("content-type", "")


def test_post_dice_export_excel_route_no_match(monkeypatch):
    service = DiceService()
    job1 = JobPosting(job_title="Cloud Engineer 1", company="CloudCorp 1", lead_source="Dice")
    service._results = [job1]

    client = _make_test_client(service, monkeypatch)
    resp = client.post("/api/dice/export/excel", json={"selected_ids": ["non-existent-id"]})
    assert resp.status_code == 400
    assert "No selected Dice job leads match" in resp.json()["detail"]
