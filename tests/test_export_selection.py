"""Test selective lead export (filtering by selected_ids) and sequence column numbering."""
import shutil
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest
from httpx import AsyncClient, ASGITransport

from app.dashboard.router import router
from app.excel.exporter import ExcelExporter
from main import app
from app.models.job import JobPosting
from app.services.scraper_service import get_scraper_service

_OUT_DIR = Path(__file__).resolve().parent / "_tmp_excel_selection"


@pytest.fixture(autouse=True)
def clean_tmp_dir():
    if _OUT_DIR.exists():
        shutil.rmtree(_OUT_DIR)
    yield
    if _OUT_DIR.exists():
        shutil.rmtree(_OUT_DIR)


def test_excel_exporter_sequence_numbers():
    jobs = [
        JobPosting(job_title="Dev 1", company="Company 1", job_url="https://example.com/1"),
        JobPosting(job_title="Dev 2", company="Company 2", job_url="https://example.com/2"),
        JobPosting(job_title="Dev 3", company="Company 3", job_url="https://example.com/3"),
    ]
    path = ExcelExporter().export(jobs, output_dir=str(_OUT_DIR), query="Dev")
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    # Check headers
    header_row = [c.value for c in ws[7]]
    assert header_row[0] == "#"
    assert header_row[1] == "Job Title"

    # Check rows sequence: 1, 2, 3
    assert ws.cell(row=8, column=1).value == 1
    assert ws.cell(row=8, column=2).value == "Dev 1"

    assert ws.cell(row=9, column=1).value == 2
    assert ws.cell(row=9, column=2).value == "Dev 2"

    assert ws.cell(row=10, column=1).value == 3
    assert ws.cell(row=10, column=2).value == "Dev 3"


@pytest.mark.asyncio
async def test_excel_export_api_filtering():
    service = get_scraper_service()
    j1 = JobPosting(job_title="Job A", company="Comp A", job_url="https://example.com/a")
    j2 = JobPosting(job_title="Job B", company="Comp B", job_url="https://example.com/b")
    j3 = JobPosting(job_title="Job C", company="Comp C", job_url="https://example.com/c")
    service._results = [j1, j2, j3]

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Export only j1 and j3 via POST
        res = await client.post("/api/export/excel", json={"selected_ids": [str(j1.id), str(j3.id)]})
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        # Export only j2 via GET query
        res_get = await client.get(f"/api/export/excel?selected_ids={j2.id}")
        assert res_get.status_code == 200


@pytest.mark.asyncio
async def test_sharepoint_export_api_filtering():
    service = get_scraper_service()
    j1 = JobPosting(job_title="Job 1", company="Comp 1")
    j2 = JobPosting(job_title="Job 2", company="Comp 2")
    service._results = [j1, j2]

    # Verify service.export_sharepoint filters by selected_ids
    # Mocking GraphSharePointExporter to inspect filtered list
    from unittest.mock import AsyncMock, patch

    with patch("app.sharepoint.graph_exporter.GraphSharePointExporter.export_jobs", new_callable=AsyncMock) as mock_export:
        mock_export.return_value = 1
        count = await service.export_sharepoint(selected_ids=[str(j1.id)])
        assert count == 1
        assert len(mock_export.call_args[0][0]) == 1
        assert mock_export.call_args[0][0][0].id == j1.id
